"""Read Microsoft Word (.docx) into a Document.

Uses python-docx. Word's named styles ("Heading 1".."Heading 6", "Title",
"List Bullet"/"List Number", "Quote") map cleanly onto our block model, which
makes Word the highest-fidelity input format.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from ..models import BlockType, ContentBlock, Document, Section
from ..utils import section_slug

_META_RE = re.compile(r"^([A-Za-z][A-Za-z0-9 _/-]{1,40}):\s*(.+)$")
# House editorial markers that authors type literally into Word.
# House heading markers authors type literally: "[H2] ", "H2: ", "H2. ".
_HEADING_MARK = re.compile(r"^(\[H[1-6]\]|H[1-6][:.])\s*", re.IGNORECASE)
# CMS scaffold lines that must never become the page title (the domain,
# "CMS STAGE …", a KV header row, "PAGE CONTENT BEGINS").
_SCAFFOLD_TITLE_RE = re.compile(
    r"^(cms stage\b|galapagosislands\.travel$|page content begins"
    r"|(publisher|author|slug|redirect|canonical|focus keyword|page template"
    r"|parent page|publish url|publish status|title tag|meta description"
    r"|secondary keywords|schema types|date (published|modified))\s*:)",
    re.IGNORECASE,
)
# A Title-Case template label ending in "Page" (e.g. "Galápagos Sea Lion Page").
# Case-sensitive so ordinary prose ending in "page" doesn't match.
_SCAFFOLD_PAGE_LABEL_RE = re.compile(r"^([A-ZÁÉÍÓÚ][\wÁÉÍÓÚáéíóúñ-]*\s+){1,4}Page$")


def _is_scaffold_title(text: str) -> bool:
    text = (text or "").strip()
    return bool(_SCAFFOLD_TITLE_RE.match(text) or _SCAFFOLD_PAGE_LABEL_RE.match(text))


# GEO/AIO snippet scaffolding written as prose: instruction lines to drop.
_GEO_SCAFFOLD_RE = re.compile(
    r"^(GEO\s*/?\s*AIO\b|AIO\b.*\bSNIPPET|Insert (this|as) .*(paragraph|answer)"
    r"|Text:\s*$)",
    re.IGNORECASE,
)
# A webmaster INSTRUCTION about the Quick Answer/GEO block (how/where to place
# it) — never the answer itself, and never body content.
_GEO_INSTRUCTION_RE = re.compile(
    r"is your GEO\b|generative engine optimization|written to be cited"
    r"|place this (block|as|snippet|text)|do not bury|very first (visible )?element"
    r"|style (this|it|as) a\b|webmaster instruction"
    r"|optimized for (ai|google|llm|voice)|first paragraph of the page"
    r"|do not use as a subtitle|\b50-word (block|citation)|do not skip it",
    re.IGNORECASE,
)


def _is_geo_instruction(rows: list[list[str]]) -> bool:
    text = " ".join(c for row in rows for c in row if c)[:400]
    return bool(_GEO_INSTRUCTION_RE.search(text))


def _looks_like_jsonld(text: str) -> bool:
    """A raw schema.org JSON-LD block (already captured into seo_schema) that must
    not leak into the page body as prose."""
    t = (text or "").lstrip()
    if t[:1] not in ("{", "["):
        return False
    return "@context" in t[:400] or "@graph" in t[:400]
_EDITORIAL_FLAG = re.compile(
    r"^(⚠️\s*)?(WEBMASTER\b|.*\bDo not publish\b|\[?VERIFY\]?\b)", re.IGNORECASE
)

# CMS-Stage annotated docs: a GEO citation-ready summary written as PROSE inside
# a "[CARD START — GEO SUMMARY …]" … "[CARD END]" wrapper (not a table). The
# inner text is the geo_answer; the markers are dropped.
_CARD_GEO_START_RE = re.compile(r"^\[CARD\s+START\b[^\]]*\bGEO\b", re.IGNORECASE)
_CARD_END_RE = re.compile(r"^\[CARD\s+END\b", re.IGNORECASE)
# Bracketed CMS annotations to drop: [SEO: …], [SCHEMA: …], [GEO: …instruction],
# [AIO …], [IMAGE …], [DATASET …], [KEYWORD …], [H2] handled elsewhere.
_BRACKET_ANNOT_RE = re.compile(
    r"^\[(SEO|SCHEMA|GEO|AIO|IMAGE|ALT|DATASET|META|KEYWORD|SECTION|NOTE|DESIGN|CARD)\b",
    re.IGNORECASE,
)
# "STAGE 8 — CMS-READY ANNOTATED CONTENT" banner.
_STAGE_SCAFFOLD_RE = re.compile(r"^STAGE\s+\d+\s*[—–-]\s*CMS", re.IGNORECASE)
# Publisher strip: "Galapagos Islands Travel | Updated 2026 | https://…".
_PUBLISHER_LINE_RE = re.compile(r"\|\s*Updated\s+\d{4}\s*\|.*https?://", re.IGNORECASE)
# A byline near the top: "By Juan Magallanes, Naturalist Expert Contributor — …".
# Allow an optional leading "the" so a corporate byline ("By the Voyagers Travel
# Company Editorial Team | GalapagosIslands.travel") is recognized, not just a
# personal one ("By Jane Darwin").
_BYLINE = re.compile(r"^By\s+(?:the\s+)?[A-Z][\w'.-]+\s+[A-Z]")

# A standalone domain/URL token within a byline (dropped from the author name).
_DOMAIN_TOKEN_RE = re.compile(
    r"^(?:https?://\S+|[\w-]+\.(?:travel|com|org|net|co|io|gov|edu)\b.*)$", re.IGNORECASE
)


def _clean_byline(text: str) -> str:
    # Strip the "By " / "By the " lead-in, then drop only the site/domain token —
    # keeping the name, role AND any company. So
    # "By Juan Magallanes, Naturalist Expert Contributor — GalapagosIslands.travel
    #  / Voyagers Travel Company" -> "Juan Magallanes, Naturalist Expert
    #  Contributor — Voyagers Travel Company", and "By the Voyagers Travel Company
    #  Editorial Team | GalapagosIslands.travel" -> "Voyagers Travel Company
    #  Editorial Team".
    s = re.sub(r"^By\s+(?:the\s+)?", "", text, flags=re.IGNORECASE).strip()
    parts = [p.strip() for p in re.split(r"\s*\|\s*|\s+[—–]\s+|\s*/\s*", s) if p.strip()]
    kept = [p for p in parts if not _DOMAIN_TOKEN_RE.match(p)]
    return " — ".join(kept) if kept else s


_BYLINE_ROLE_RE = re.compile(
    r"contributor|naturalist|editor|writer|correspondent|journalist|guide", re.IGNORECASE
)


def _looks_like_byline(text: str) -> bool:
    """A standalone author byline ('By <Name>, <Role> — <site>'), as opposed to a
    body sentence that merely starts with 'By'. Requires a real byline signal
    (an em/en dash, a site domain, or a role word) so prose like 'By Darwin's
    account…' is never mistaken for one."""
    if not _BYLINE.match(text) or len(text) > 220:
        return False
    low = text.lower()
    return (
        "—" in text or "–" in text or ".travel" in low or ".com" in low
        or bool(_BYLINE_ROLE_RE.search(low))
    )


# Strong "internal / do-not-publish" markers. Scanned across the WHOLE table
# because these blocks often lead with a separator rule, pushing the real marker
# into a later row. Kept CTA-agnostic (no 'Voyagers'/'Latin Trails') so real CTA
# tables are never dropped.
_INSTR_HEAD_CHARS = ("⚠", "▶", "■", "⬛", "⚑", "🟩", "🟢", "🛑")
_INSTR_MARKERS = (
    "do not publish", "internal cms", "cms use only", "internal use only",
    "cms editor", "publisher header", "webmaster", "green box",
    "internal instruction", "requires editorial review", "version footer",
    "verify items", "access note", "for internal use",
)


def _is_instruction_table(rows: list[list[str]]) -> bool:
    """Internal-instruction / explainer tables (green boxes) are not content."""
    if not rows:
        return False
    head = " ".join(rows[0]).strip()
    full = " ".join(c for row in rows for c in row).lower()
    looks_internal = (
        head.startswith(_INSTR_HEAD_CHARS)
        or "what this is" in full
        or any(m in full for m in _INSTR_MARKERS)
    )
    if not looks_internal:
        return False
    # Some docs pack the real dual-CTA inside the same green box as the notes.
    # Keep those so the CTA path (which dedupes to 2) can mine them.
    if _AUD_SPLIT_RE.search(full) or re.search(r"\bcta\s*\d", full):
        return False
    return True


_SCAFFOLD_RE = re.compile(r"(?i)^(publish at|publish url|place under|slug\b|version\b|do not publish)")


def _table_target(table, body_order: dict, section_start_pos: list, doc, current):
    """The section a stray table belongs to: the one whose heading most recently
    precedes it in document order (not just the last section parsed)."""
    tpos = body_order.get(table._tbl, 1 << 30)
    target = current if not doc.sections else doc.sections[-1]
    best_pos = -2
    for pos, sec in section_start_pos:
        if pos <= tpos and pos > best_pos:
            best_pos, target = pos, sec
    return target


def _is_sources_table(rows: list[list[str]]) -> bool:
    """A 2-column citations table whose rows carry URLs."""
    if _ncols(rows) < 2 or len(rows) < 2:
        return False
    with_urls = sum(1 for r in rows if _URL_RE.search(" ".join(r)))
    return with_urls >= max(2, len(rows) // 2)


def _extract_sources_table(rows: list[list[str]]) -> list[dict]:
    out: list[dict] = []
    for r in rows:
        m = _URL_RE.search(" ".join(r))
        if not m:
            continue
        url = m.group(0).rstrip(".,);")
        label = r[0].strip().rstrip("—–-. ").strip()
        if label:
            out.append({"label": label, "url": url})
    return out


# A GEO/AIO answer block, in any of its house variants. The answer is the prose
# AFTER the marker line ("AIO / GEO BLOCK …", "PUBLISH THIS ANSWER TEXT:", etc.).
_GEO_MARKER_RE = re.compile(
    r"(AIO\b|GEO\s+(?:BLOCK|ANSWER)|PUBLISH THIS ANSWER TEXT|answer extraction|Quick Answer)",
    re.IGNORECASE,
)
_PUBLISH_ANSWER_RE = re.compile(r"PUBLISH(?:ABLE)?[^\n:]*ANSWER[^\n:]*:\s*", re.IGNORECASE)
_GEO_STOP_RE = re.compile(r"^(PLACE UNDER|PUBLISH URL|SLUG\b|■|⚠|⬛|VERSION)", re.IGNORECASE)
_GEO_EXPLAINER_RE = re.compile(r"^(WHAT THIS|WHY THIS|HOW \b|HOW THIS)", re.IGNORECASE)


def _extract_geo_answer(rows: list[list[str]]) -> str:
    """Pull the publishable ~50-word GEO/AI answer out of a GEO block table.

    Handles two house styles: an explicit 'PUBLISH THIS ANSWER TEXT:' marker
    (answer follows it), or an 'AIO / GEO …' header whose answer is the first
    prose after the marker line (skipping any 'WHAT THIS IS' explainer).
    """
    text = "\n".join(cell for row in rows for cell in row if cell).strip()
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]

    # Case 1: explicit publishable-answer marker.
    for idx, line in enumerate(lines):
        if _PUBLISH_ANSWER_RE.search(line):
            rest = _PUBLISH_ANSWER_RE.sub("", line).strip()
            collected = [rest] if rest else []
            for nxt in lines[idx + 1:]:
                if _GEO_STOP_RE.match(nxt):
                    break
                collected.append(nxt)
            ans = " ".join(" ".join(collected).split()).strip()
            if ans:
                return ans

    # Case 2: AIO / GEO / "Quick Answer" header, answer is the first prose after
    # the marker. The marker may sit on its own line (answer follows) or inline
    # with the answer on the same line ("QUICK ANSWER: <answer>").
    if not _GEO_MARKER_RE.search(text[:160]):
        return ""
    out: list[str] = []
    started = False
    for line in lines:
        if not started:
            m = _GEO_MARKER_RE.search(line)
            if not m:
                continue
            started = True
            # Inline answer only when the marker is directly followed by
            # ": <prose>" (e.g. "QUICK ANSWER: Baltra …"), not when trailing
            # header tokens follow (e.g. "AIO / GEO BLOCK …").
            inline = re.match(r"\s*[:–—-]\s*(.+)", line[m.end():])
            if inline:
                rest = inline.group(1).strip()
                if (
                    len(rest.split()) >= 6
                    and not _GEO_EXPLAINER_RE.match(rest)
                    and not _GEO_STOP_RE.match(rest)
                ):
                    out.append(rest)
            continue
        if _GEO_STOP_RE.match(line):
            break
        if _GEO_EXPLAINER_RE.match(line):
            if out:
                break
            continue  # skip explainer text before the answer
        out.append(line)
    return " ".join(" ".join(out).split()).strip()


# --- structured tables ----------------------------------------------------- #
def _ncols(rows: list[list[str]]) -> int:
    return max((len(r) for r in rows), default=0)


def _is_visitor_sites_table(rows: list[list[str]]) -> bool:
    head = " ".join(c.lower() for c in rows[0])
    return _ncols(rows) >= 3 and ("visitor site" in head or ("site" in head and "access" in head))


def _is_quick_facts_table(rows: list[list[str]]) -> bool:
    """A 2-column label|value reference table (short labels, no prose header)."""
    if _ncols(rows) != 2 or len(rows) < 3:
        return False
    labels = [r[0].strip() for r in rows if len(r) >= 2 and r[0].strip()]
    return bool(labels) and all(len(label) <= 40 for label in labels)


# A seasonal wildlife/what-to-see calendar: 2 columns, a "season/period/month"
# header on the left and a "wildlife/highlights/conditions" header on the right.
_CAL_SEASON_HEAD = ("season", "when to", "time of year")
_CAL_MONTH_HEAD = ("month", "period", "dates")
_CAL_HL_HEAD = (
    "wildlife", "highlight", "condition", "what to see", "activity",
    "to see", "prioritize", "what to",
)


def _is_wildlife_calendar_table(rows: list[list[str]]) -> bool:
    if _ncols(rows) < 2 or len(rows) < 2:
        return False
    h0 = rows[0][0].lower()
    rest = " ".join(rows[0][1:]).lower()
    has_season = any(k in h0 for k in _CAL_SEASON_HEAD) or any(k in h0 for k in _CAL_MONTH_HEAD)
    return has_season and any(k in rest for k in _CAL_HL_HEAD)


def _norm_ws(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\n", " ")).strip()


def _extract_wildlife_calendar(rows: list[list[str]]) -> list[dict]:
    """Season → highlights, column-aware. Handles both a 2-column
    'Season | Wildlife Highlights' table and a wider 'Season | Months |
    Conditions | What to Prioritize' one, mapping period/label/highlights by
    their header. A parenthetical in the season cell becomes the label."""
    header = [c.lower() for c in rows[0]]
    season_idx = [i for i, h in enumerate(header) if any(k in h for k in _CAL_SEASON_HEAD)]
    month_idx = [i for i, h in enumerate(header) if any(k in h for k in _CAL_MONTH_HEAD)]
    hl_idx = [i for i, h in enumerate(header) if any(k in h for k in _CAL_HL_HEAD)]
    period_cols = month_idx or season_idx or [0]
    label_cols = season_idx if month_idx else []
    if not hl_idx:
        hl_idx = [i for i in range(1, len(header)) if i not in period_cols and i not in label_cols]

    def cell(row: list[str], i: int) -> str:
        return (row[i] or "").strip() if i < len(row) else ""

    raw_header = rows[0]
    multi = len(hl_idx) > 1
    out: list[dict] = []
    for r in rows[1:]:
        period = " ".join(v for i in period_cols if (v := cell(r, i)))
        label = " ".join(v for i in label_cols if (v := cell(r, i)))
        m = re.search(r"\(([^)]*)\)", period)
        if m and not label:
            label = m.group(1).strip()
            period = (period[: m.start()] + period[m.end():]).strip()
        # With more than one highlight column (e.g. "Conditions" + "What to
        # Prioritize") label each by its header so they don't run together; a
        # single column stays as plain prose.
        hl_parts: list[str] = []
        for i in hl_idx:
            v = _norm_ws(re.sub(r"\[VERIFY[^\]]*\]", "", cell(r, i)))
            if not v:
                continue
            head = _norm_ws(raw_header[i] if i < len(raw_header) else "").rstrip(":")
            hl_parts.append(f"**{head}:** {v}" if multi and head else v)
        highlights = "\n\n".join(hl_parts)
        period, label = _norm_ws(period).strip(" —–-"), _norm_ws(label)
        if period or highlights:
            out.append({"period": period, "label": label, "highlights": highlights})
    return out


def _infer_access_type(access: str) -> str:
    a = access.lower()
    if "cruise" in a:
        return "Cruise-only"
    if any(w in a for w in ("walk", "taxi", "bus", "free", "drive", "ferry", "road", "town")):
        return "Land-based"
    return ""


def _extract_visitor_sites(rows: list[list[str]]) -> list[dict]:
    header = [c.lower() for c in rows[0]]

    def col(*names: str) -> int | None:
        for i, h in enumerate(header):
            if any(n in h for n in names):
                return i
        return None

    i_name, i_access = col("site"), col("access")
    i_wild, i_notes = col("wildlife", "species"), col("note", "description")

    def cell(row: list[str], i: int | None) -> str:
        return row[i].strip() if i is not None and i < len(row) else ""

    out: list[dict] = []
    for r in rows[1:]:
        name = cell(r, i_name)
        if not name:
            continue
        access = cell(r, i_access)
        out.append({
            "site_name": name,
            "access": access,
            "access_type": _infer_access_type(access),
            "species_seen": cell(r, i_wild),
            "description": cell(r, i_notes),
        })
    return out


def _extract_quick_facts(rows: list[list[str]]) -> list[dict]:
    start = 1 if {c.lower() for c in rows[0]} & {"fact", "detail", "label", "value", "item"} else 0
    out: list[dict] = []
    for r in rows[start:]:
        if len(r) >= 2 and r[0].strip() and r[1].strip():
            out.append({"label": r[0].strip(), "value": r[1].strip()})
    return out


# ── Wildlife species pages ──────────────────────────────────────────────────
# The intended-page-type hints that mean "single wildlife species page", so the
# reader keeps a table-heavy species doc out of the CMS-Stage adapter.
_SPECIES_PAGE_TYPES = {
    "wildlife_single", "wildlife_species", "species_page", "wildlife_page",
}
# Informative/guide pages have no wildlife/island structured fields (seasonality,
# subspecies, quick_facts, visitor_sites, wildlife_calendar). Their multi-column
# tables are ordinary DATA tables that must survive as HTML in the section
# content — so those structured-table detectors are skipped for this page type,
# otherwise a "Month | …" or "…| Best Season |…" table gets mined into metadata
# the informative profile never reads and vanishes from the page.
_INFORMATIVE_PAGE_TYPES = {
    "informative", "informative_page", "planning", "info",
}
# Content signals (for preview, when no page type is passed). Kept specific so
# island docs — which merely LINK to /wildlife/ pages — never match.
_SPECIES_SIGNAL_RE = re.compile(
    r"taxonconcept|species\s+guide|scientific\s*name\s*[:|]|\bscientificname\b",
    re.IGNORECASE,
)


def _norm_page_type(value: str | None) -> str:
    return (value or "").strip().lower().replace(" ", "_")


def _looks_like_species(full_text: str) -> bool:
    return bool(_SPECIES_SIGNAL_RE.search(full_text or ""))


# Publisher / SEO scaffold KV tables (PUBLISH URL, Title tag, Anchor Text…) that
# a species doc packs into 2-column tables. They look like a quick-facts table
# but are webmaster metadata, not content — skip them so they don't pollute the
# "At a Glance" repeater.
_SCAFFOLD_LABELS = {
    "publish url", "parent page", "page template", "redirect", "publisher",
    "author / byline", "author", "date published", "date modified", "title tag",
    "meta description", "canonical", "canonical url", "focus keyword",
    "secondary keywords", "schema types", "anchor text", "target url",
    "location", "publish status", "slug", "url",
}


def _scaffold_label(label: str) -> str:
    return re.sub(r"\s*\(.*?\)\s*", " ", (label or "").lower()).strip().rstrip(":").strip()


_AUTHOR_KV_LABELS = {"author", "author / byline", "author/byline", "byline", "by"}


def _author_from_kv(rows: list[list[str]]) -> str:
    """Pull an 'Author / Byline' value out of a webmaster header KV table."""
    for r in rows:
        if len(r) >= 2 and _scaffold_label(r[0]) in _AUTHOR_KV_LABELS:
            val = re.sub(r"^By\s+", "", r[1].strip()).strip()
            val = re.split(r"\s+[—–]\s+", val, maxsplit=1)[0].strip()
            if val:
                return val
    return ""


def _is_scaffold_kv_table(rows: list[list[str]]) -> bool:
    """A 2-col table whose row labels are mostly webmaster/SEO scaffold keys."""
    if _ncols(rows) != 2 or len(rows) < 2:
        return False
    labels = [_scaffold_label(r[0]) for r in rows if len(r) >= 2 and r[0].strip()]
    if not labels:
        return False
    hits = sum(1 for lbl in labels if lbl in _SCAFFOLD_LABELS)
    return hits >= max(2, (len(labels) + 1) // 2)


# Dedicated species-fact labels mined out of the "At a Glance" facts table.
_FACT_KEYS = {
    "scientific_name": ("scientific name", "scientific"),
    "common_name": ("common name",),
    "conservation_status": ("iucn status", "iucn", "conservation status", "status"),
    "population": ("population estimate", "population", "world population"),
}
# IUCN Red List categories, longest-first so "Critically Endangered" wins over
# "Endangered". Maps free text onto the ACF select's canonical choices.
_IUCN_CHOICES = [
    "Critically Endangered", "Near Threatened", "Least Concern", "Data Deficient",
    "Vulnerable", "Endangered",
]


def _normalize_iucn(value: str) -> str:
    low = (value or "").lower()
    for choice in _IUCN_CHOICES:
        if choice.lower() in low:
            return choice
    return ""


def _extract_species_facts(quick_facts: list[dict]) -> dict:
    """Pull scientific/common name, IUCN status and population out of the
    already-parsed facts rows. Returns only the fields actually found."""
    out: dict = {}
    for row in quick_facts:
        label = _scaffold_label(row.get("label", ""))
        value = (row.get("value") or "").strip()
        if not label or not value:
            continue
        for field, keys in _FACT_KEYS.items():
            if field in out:
                continue
            if any(label == k or label.startswith(k) for k in keys):
                if field == "conservation_status":
                    norm = _normalize_iucn(value)
                    if norm:
                        out[field] = norm
                else:
                    out[field] = value
                break
    return out


# Seasonality calendar: a Period/Month table (period + optional status + notes).
_SEASON_PERIOD_HEADS = ("month", "period", "season", "dates", "when", "time of year")
_SEASON_LABEL_HEADS = ("status", "label", "stage", "phase", "presence")


def _is_seasonality_table(rows: list[list[str]]) -> bool:
    if _ncols(rows) < 2 or len(rows) < 3:
        return False
    h0 = rows[0][0].lower()
    return any(k in h0 for k in _SEASON_PERIOD_HEADS)


def _extract_seasonality(rows: list[list[str]]) -> list[dict]:
    header = [c.lower() for c in rows[0]]

    def find(cands) -> int | None:
        for i, h in enumerate(header):
            if any(c in h for c in cands):
                return i
        return None

    i_period = find(_SEASON_PERIOD_HEADS) or 0
    i_label = find(_SEASON_LABEL_HEADS)
    # Notes = the first remaining column (Event / Notes / Highlights / …).
    i_notes = next(
        (i for i in range(len(header)) if i not in (i_period, i_label)), None
    )

    def cell(r: list[str], i: int | None) -> str:
        return _norm_ws(r[i]) if i is not None and i < len(r) else ""

    out: list[dict] = []
    for r in rows[1:]:
        period = cell(r, i_period)
        if not period:
            continue
        out.append({
            "period": period,
            "label": cell(r, i_label),
            "notes": cell(r, i_notes),
        })
    return out


# Subspecies / island-variants table: an "Island" column plus a name column
# (Species / Subspecies / Name) — e.g. giant tortoise shell types, marine iguana
# subspecies.
_SUB_NAME_HEADS = ("subspecies", "species", "name", "race", "form", "variant")


def _is_subspecies_table(rows: list[list[str]]) -> bool:
    if _ncols(rows) < 2 or len(rows) < 3:
        return False
    header = [c.lower() for c in rows[0]]
    if "island" not in header[0]:
        return False
    return any(any(k in h for k in _SUB_NAME_HEADS) for h in header[1:])


def _extract_subspecies(rows: list[list[str]]) -> list[dict]:
    header = [c.lower() for c in rows[0]]

    def find(cands, skip=()) -> int | None:
        for i, h in enumerate(header):
            if i in skip:
                continue
            if any(c in h for c in cands):
                return i
        return None

    i_island = 0
    i_name = find(_SUB_NAME_HEADS, skip={0})
    i_pop = find(("population", "approx"), skip={0})
    i_status = find(("status",), skip={0})
    used = {i_island, i_name, i_pop, i_status}
    i_trait = find(("shell", "trait", "notable", "distinguishing", "descr", "feature"),
                   skip=used) or next(
        (i for i in range(1, len(header)) if i not in used), None)

    def cell(r: list[str], i: int | None) -> str:
        return _norm_ws(r[i]) if i is not None and i < len(r) else ""

    out: list[dict] = []
    for r in rows[1:]:
        island = cell(r, i_island)
        name = cell(r, i_name)
        if not island and not name:
            continue
        out.append({
            "island": island,
            "name": name,
            "trait": cell(r, i_trait),
            "population": cell(r, i_pop),
            "status": cell(r, i_status),
        })
    return out


_URL_RE = re.compile(r"https?://\S+")
_TRADE_RE = re.compile(r"\b(trade|latin trails|dmc|wholesaler|agent|group programs)\b", re.IGNORECASE)
_CTA_BUTTON_RE = re.compile(r'CTA button:\s*"(.+?)"\s*(?:→|->)\s*(\S+)', re.IGNORECASE)
# Visible text inside a paragraph, INCLUDING content controls (w:sdt) and
# hyperlinks, which python-docx's `.text` silently drops.
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_WT = _W_NS + "t"
_WBR = _W_NS + "br"  # soft line break (Shift+Enter)
_WCR = _W_NS + "cr"  # carriage return
_WTAB = _W_NS + "tab"
_WHYPERLINK = _W_NS + "hyperlink"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_INTERNAL_LINK_RE = re.compile(r"\[INTERNAL LINK:\s*(.+?)\s*(?:→|->)\s*([^\]\s]+)\s*\]")
# "⬛ WEBMASTER DESIGN NOTE: …" styling instructions — never published, even when
# appended inline to a real paragraph. Stripped from the marker to end of line.
_DESIGN_NOTE_RE = re.compile(
    r"\s*[^\w\s\n]*\s*WEBMASTER\s+DESIGN\s+NOTE\b[^\n]*", re.IGNORECASE
)


def _run_text(el) -> str:
    """Visible text under a run-level element (soft breaks/tabs preserved)."""
    parts: list[str] = []
    for node in el.iter():
        if node.tag == _WT:
            parts.append(node.text or "")
        elif node.tag in (_WBR, _WCR):
            parts.append("\n")
        elif node.tag == _WTAB:
            parts.append("\t")
    return "".join(parts)


def _p_text(paragraph) -> str:
    # Walk direct children in document order so soft line breaks (<w:br>) and tabs
    # survive as real separators, and so a <w:hyperlink> becomes a real Markdown
    # link ([label](url)) — python-docx's `.text` drops the URL entirely, which
    # loses every internal link (e.g. the "Explore More" SEO footer).
    try:
        rels = paragraph.part.rels
    except Exception:  # pragma: no cover - a detached paragraph has no part
        rels = {}
    parts: list[str] = []
    for child in paragraph._p:
        if child.tag == _WHYPERLINK:
            label = _run_text(child)
            rid = child.get(_R_ID)
            url = ""
            if rid and rid in rels:
                try:
                    url = rels[rid].target_ref or ""
                except Exception:  # pragma: no cover
                    url = ""
            parts.append(f"[{label}]({url})" if (url and label) else label)
        else:
            parts.append(_run_text(child))
    return "".join(parts)


def _cell_text(cell) -> str:
    return "\n".join(_p_text(p) for p in cell.paragraphs).strip()


# A visual rule made of box-drawing / dash / underscore / equals runs. These are
# layout separators in the house docs, never content — drop them so they don't
# become empty feature sections or trail into the previous card's prose.
_SEP_CHARS = set("─━╌╍—–-_=~⎯。・ \t")
_SEP_KEY = set("─━╌╍—–-_=")


def _is_separator(text: str) -> bool:
    t = text.strip()
    if len(t) < 3 or any(ch.isalnum() for ch in t):
        return False
    return all(ch in _SEP_CHARS for ch in t) and any(ch in _SEP_KEY for ch in t)


# "→ INTERNAL LINK: /url/ [| See also: /url2/ | /url3/]" — the arrow-style
# internal-link notation used across the house docs (no brackets, no label).
_ARROW_LINK_RE = re.compile(r"(?:→|->)\s*INTERNAL LINK:\s*([^\n]+)", re.IGNORECASE)
_ANY_URL_RE = re.compile(r"https?://\S+|/[\w\-./]+/?")


def _slug_label(url: str) -> str:
    """Derive a human label from a URL slug: /wildlife/giant-tortoise/ -> 'Giant Tortoise'."""
    slug = url.rstrip("/").rsplit("/", 1)[-1]
    slug = re.sub(r"[-_]+", " ", slug).strip()
    return slug.title() if slug else "Learn more"


def _arrow_links_to_md(text: str) -> str:
    """'→ INTERNAL LINK: /url/ | …' -> Markdown link(s) with slug-derived labels.

    Drops the arrow + "INTERNAL LINK:"/"See also:" scaffolding, keeping each URL as
    a real link. A standalone link paragraph later becomes a card button; an inline
    one stays an inline <a>.
    """
    def repl(m: re.Match) -> str:
        urls = _ANY_URL_RE.findall(m.group(1))
        if not urls:
            return ""
        return " ".join(f"[{_slug_label(u)}]({u})" for u in urls)

    return _ARROW_LINK_RE.sub(repl, text)


def _clean_markers(text: str) -> str:
    """[INTERNAL LINK: label → /url/] -> Markdown [label](/url/).

    WYSIWYG fields render the Markdown link to a real <a>; textarea fields strip
    it back to the label (they can't hold links). Either way the URL is kept here.
    """
    text = _DESIGN_NOTE_RE.sub("", text)
    # Drop inline editorial "verify this fact" notes, e.g. "[VERIFY 2026 fee]".
    # These are internal flags and must never reach the published page.
    text = _INLINE_VERIFY_RE.sub("", text)
    text = _INTERNAL_LINK_RE.sub(r"[\1](\2)", text)
    return _arrow_links_to_md(text)


# Inline "[VERIFY …]" editorial note anywhere in a line (with any leading space).
_INLINE_VERIFY_RE = re.compile(r"\s*\[VERIFY[^\]]*\]", re.IGNORECASE)


_CTA_SIGNALS = (
    "cta button:", "call to action", "dual cta block", "with a specialist",
    "book with voyagers", "direct travelers — book", "direct travellers — book",
    "voyagers travel company", "latin trails", "planning your galápagos",
    "planning your galapagos", "independent travellers", "independent travelers",
)


def _is_cta_table(rows: list[list[str]]) -> bool:
    text = " ".join(c for row in rows for c in row).lower()
    return any(sig in text for sig in _CTA_SIGNALS)


_TRADE_MARK_RE = re.compile(
    r"latin trails|dmc|trade use|travel agent|tour operator|wholesal|"
    r"industry partner|group program|travel trade",
    re.IGNORECASE,
)
# Audience label ("DIRECT TRAVELERS", "TRAVEL TRADE"…) up to its separator, which
# may be a colon OR an em/en dash ("DIRECT TRAVELERS — Book with…").
_AUD_SPLIT_RE = re.compile(
    r"(INDEPENDENT TRAVELLERS?|DIRECT TRAVEL\w*|TRAVEL AGENTS?|TRAVEL TRADE|TRADE\b)"
    r"[^\n:—–]*[:—–]",
    re.IGNORECASE,
)
_FIELD_PREFIX_RE = re.compile(r"^(website|email|web|tel|phone|button)\s*:\s*", re.IGNORECASE)


def _cta_audience(text: str) -> str:
    return "Travel trade" if _TRADE_MARK_RE.search(text) else "Direct travelers"


def _split_audience_segments(cell: str) -> list[str]:
    marks = list(_AUD_SPLIT_RE.finditer(cell))
    if len(marks) < 2:
        return []
    return [
        cell[m.start(): (marks[i + 1].start() if i + 1 < len(marks) else len(cell))].strip()
        for i, m in enumerate(marks)
    ]


# A button line like "Contact Latin Trails → https://…" or "Book now → /cruises/".
_CONTACT_BTN_RE = re.compile(r"^(.{0,60}?)\s*(?:→|->)\s*(.+)$")
# A bare domain (no scheme), e.g. "galapagosislands.travel/contact", "latintrails.com".
_BARE_DOMAIN_RE = re.compile(
    r"\b([a-z0-9][a-z0-9-]*(?:\.[a-z0-9-]+)*\.(?:com|travel|org|net|io|co)(?:\.[a-z]{2})?(?:/\S*)?)",
    re.IGNORECASE,
)


def _norm_url(u: str) -> str:
    u = u.rstrip(".,);")
    if u.startswith(("http://", "https://", "/", "mailto:")):
        return u
    return "https://" + u


_CTA_HEADER_RE = re.compile(r"^\s*CTA\s*[:#-]?\s*[^\n]*\n", re.IGNORECASE)


def _cta_from_text(cell: str) -> dict:
    # Single-cell "CTA: <audience> | <company>" header line (giant-tortoise style):
    # not the headline — derive audience from it, then drop it so the real hook
    # becomes the title.
    forced_aud = ""
    hdr = _CTA_HEADER_RE.match(cell)
    if hdr:
        forced_aud = cell[: hdr.end()]
        cell = cell[hdr.end():]
    btn = _CTA_BUTTON_RE.search(cell)
    cell_wo = _CTA_BUTTON_RE.sub("", cell)
    # House format: "INDEPENDENT TRAVELLERS: <body…>" — the audience label is a
    # prefix, not a headline; everything after it is body copy (no separate title).
    aud_prefix = _AUD_SPLIT_RE.match(cell_wo.strip())
    if aud_prefix:
        cell_wo = cell_wo[aud_prefix.end():].strip()
    lines = [ln.strip() for ln in cell_wo.split("\n") if ln.strip()]
    # After the audience label: if more copy follows, the first line is the
    # headline ("DIRECT TRAVELERS — Book with Voyagers…"); a lone line is body.
    if len(lines) > 1:
        title, rest = lines[0], lines[1:]
    else:
        title, rest = "", lines
    url = btn.group(2).strip() if btn else ""
    label = btn.group(1).strip().rstrip("→ ").strip() if btn else ""
    body: list[str] = []
    for ln in rest:
        # A "Label → url" contact/button line: lift it out of the body.
        bm = _CONTACT_BTN_RE.match(ln)
        if not btn and bm and _ANY_URL_RE.search(bm.group(2)):
            urls = _ANY_URL_RE.findall(bm.group(2))
            url = url or _norm_url(urls[-1])  # prefer the absolute (last) URL
            label = label or bm.group(1).strip()
            continue
        # A "→ <call to action>" line with no URL is the button LABEL (the URL
        # was given separately, e.g. a bare domain on another line) — not body.
        am = re.match(r"^\s*(?:→|->)\s*(.+)$", ln)
        if am and not _ANY_URL_RE.search(ln):
            label = label or am.group(1).strip()
            continue
        # A trailing bare-domain URL (no scheme, no arrow): pull it as the button,
        # but only when it sits at the END of the line — not when it opens a
        # sentence ("GalapagosIslands.travel is maintained by…").
        if not url and not _URL_RE.search(ln):
            dm = _BARE_DOMAIN_RE.search(ln)
            if dm and dm.end() >= len(ln.rstrip()):
                url = _norm_url(dm.group(1))
                ln = re.sub(r"\s*(?:→|->|—|–|-)\s*$", "", ln[: dm.start()].strip()).strip()
                if not ln:
                    continue
        if not url:
            m = _URL_RE.search(ln)
            if m:
                url = m.group(0).rstrip(".,);")
        if _FIELD_PREFIX_RE.match(ln):
            continue  # drop "Website:/Email:" label lines from body
        body.append(ln)
    block = {
        "audience": _cta_audience(forced_aud or cell),
        "title": title,
        "text": " ".join(" ".join(body).split()),
    }
    if label:
        block["button_label"] = label
    if url:
        block["button_url"] = url
    return block


# "CALL TO ACTION … CTA 1 — <audience>: Headline: … Body: … Button: label → url".
_CTA_INSTR_RE = re.compile(
    r"CTA\s*\d+\s*[—–-]\s*([^\n:]+):(.*?)(?=CTA\s*\d+\s*[—–-]|\Z)", re.IGNORECASE | re.DOTALL
)
_CTA_TRADE_RE = re.compile(r"trade|industry|partner|wholesal|agent|dmc", re.IGNORECASE)


def _grab_field(block: str, name: str) -> str:
    m = re.search(rf"{name}:\s*(.+)", block, re.IGNORECASE)
    return m.group(1).strip() if m else ""


def _parse_cta_instruction(cell: str) -> list[dict]:
    out: list[dict] = []
    for m in _CTA_INSTR_RE.finditer(cell):
        audience_text, body_block = m.group(1).strip(), m.group(2)
        block = {
            "audience": "Travel trade" if _CTA_TRADE_RE.search(audience_text) else "Direct travelers",
            "title": _grab_field(body_block, "Headline"),
            "text": _grab_field(body_block, "Body"),
        }
        bm = re.search(r"(.+?)\s*(?:→|->)\s*(.+)$", _grab_field(body_block, "Button"))
        if bm:
            block["button_label"] = bm.group(1).strip()
            url = bm.group(2).strip()
            if url.startswith(("/", "http", "mailto:")):  # ignore non-URL placeholders
                block["button_url"] = url
        out.append(block)
    return out


# Plain-paragraph CTAs (waved-albatross style): the two formal CTAs written as
# body paragraphs — "BOOK WITH VOYAGERS TRAVEL COMPANY — …", "TRAVEL INDUSTRY
# PARTNERS (agents…) — …" — instead of a CTA table. A strong leading label
# followed by an em/en dash marks one; anything softer is left as prose.
_PLAIN_CTA_LEAD_RE = re.compile(
    r"^\s*("
    r"BOOK WITH\b[^\n—–]*|"
    r"TRAVEL INDUSTRY PARTNERS\b[^\n—–]*|"
    r"TRAVEL TRADE\b[^\n—–]*|"
    r"TRAVEL AGENTS?\b[^\n—–]*|"
    r"(?:FOR\s+)?DIRECT TRAVEL\w*\b[^\n—–]*|"
    r"INDEPENDENT TRAVEL\w*\b[^\n—–]*"
    r")\s*[—–]\s+",
    re.IGNORECASE,
)


def _looks_like_plain_cta(text: str) -> bool:
    return bool(_PLAIN_CTA_LEAD_RE.match(text or ""))


def _cta_from_paragraph(text: str) -> dict | None:
    """Turn a single plain-paragraph CTA into a cta_block. The lead label
    (before the dash) becomes the title; the rest is body; a URL or bare domain
    anywhere in the body becomes the button."""
    m = _PLAIN_CTA_LEAD_RE.match(text)
    if not m:
        return None
    label = m.group(1).strip(" —–-:")
    body = text[m.end():].strip()
    # Title = the label, minus a trailing "(agents, wholesalers…)" audience note,
    # tidied from ALL-CAPS to Title Case.
    title = re.sub(r"\s*\([^)]*\)\s*$", "", label).strip()
    if title == title.upper():
        title = title.title()
    url = ""
    um = _URL_RE.search(body)
    if um:
        url = um.group(0).rstrip(".,);")
    else:
        # Only a bare domain that carries a PATH ("site.travel/contact/") is a
        # deliberate link; a bare site name dropped mid-sentence ("…or
        # GalapagosIslands.travel.") is just a mention, not a button target.
        dm = _BARE_DOMAIN_RE.search(body)
        if dm and "/" in dm.group(1):
            url = _norm_url(dm.group(1).rstrip(".,);"))
    block = {
        "audience": _cta_audience(text),
        "title": title,
        "text": " ".join(body.split()),
    }
    if url:
        block["button_url"] = url
        block["button_label"] = "Contact us" if "/contact" in url else "Learn more"
    return block


def _looks_like_cta_grid(rows: list[list[str]]) -> bool:
    """A column-organized CTA table: row 0 names each audience across columns
    (e.g. 'For Travelers Booking Direct' | 'For Travel Trade & Wholesalers'),
    and each column beneath is one CTA (company title, then body)."""
    if _ncols(rows) < 2 or len(rows) < 2:
        return False
    aud_words = (
        "traveler", "traveller", "direct", "trade", "wholesal", "agent",
        "operator", "booking",
    )
    header = [c.lower() for c in rows[0]]
    return sum(1 for h in header if any(w in h for w in aud_words)) >= 2


def _extract_cta_blocks(rows: list[list[str]]) -> list[dict]:
    """Parse the two formal CTAs (Voyagers + Latin Trails) across house formats.

    The whole table is joined first, because a CTA is often spread across cells:
    the audience+headline ("DIRECT TRAVELERS — Book with Voyagers…") in one cell,
    the body and contact lines in others. Joining lets the headline become the
    CTA title instead of being mistaken for body copy.
    """
    # Column-organized grid: row 0 = audience headers, each column below is a CTA
    # (title = the company line, text = the copy under it). A row-major join would
    # interleave the two CTAs, so handle columns explicitly.
    if _looks_like_cta_grid(rows):
        grid: list[dict] = []
        for c in range(_ncols(rows)):
            cells = [r[c].strip() for r in rows if c < len(r) and r[c].strip()]
            if len(cells) < 2:
                continue
            b = _cta_from_text("\n".join(cells[1:]))  # skip the audience header
            b["audience"] = _cta_audience(" ".join(cells))
            grid.append(b)
        grid = [b for b in grid if _cta_is_real(b)]
        if len(grid) >= 2:
            return grid
    joined = "\n".join(c.strip() for row in rows for c in row if c.strip())
    if re.search(r"call to action|CTA\s*\d+\s*[—–-]", joined, re.IGNORECASE):
        blocks = _parse_cta_instruction(joined)
        if blocks:
            return blocks
    segments = _split_audience_segments(joined)  # "INDEPENDENT…:" / "DIRECT… —" …
    if len(segments) >= 2:
        return [_cta_from_text(seg) for seg in segments]
    if _AUD_SPLIT_RE.match(joined):  # a single-audience CTA table (e.g. Floreana)
        return [_cta_from_text(joined)]
    # Fallback: simple tables with one contact/CTA block per cell.
    out: list[dict] = []
    for row in rows:
        for cell in row:
            if cell.strip():
                out.append(_cta_from_text(cell.strip()))
    return out


_CTA_KEYWORDS = (
    "book", "contact", "voyager", "latin trails", "travel company", "dmc",
    "itinerary", "plan your", "specialist", "enquir", "inquir", "trade",
)
# SEO / editorial scaffolding that sometimes carries a URL but is never a CTA.
_CTA_JUNK = (
    "title tag", "meta description", "focus keyword", "canonical", "url slug",
    "slug:", "h1 tag", "schema", "og:", "alt text", "internal link",
)


def _cta_is_real(b: dict) -> bool:
    """Reject scaffolding rows (separators, publisher/SEO notes) posing as CTAs."""
    blob = " ".join(
        [b.get("title") or "", b.get("text") or "", b.get("button_label") or ""]
    ).strip().lower()
    if any(m in blob for m in _INSTR_MARKERS) or any(m in blob for m in _CTA_JUNK):
        return False
    if b.get("button_url"):
        return True
    if not blob or _is_separator(blob):
        return False
    return any(k in blob for k in _CTA_KEYWORDS)


def _dedupe_ctas(blocks: list[dict]) -> list[dict]:
    """Reduce the raw CTA candidates to the two canonical ones (best Direct +
    best Trade), dropping junk. A block with a real button and more copy wins."""
    best: dict[str, tuple] = {}
    for b in blocks:
        if not _cta_is_real(b):
            continue
        aud = b.get("audience") or "Direct travelers"
        score = (1 if b.get("button_url") else 0, len(b.get("text") or "") + len(b.get("title") or ""))
        if aud not in best or score > best[aud][0]:
            best[aud] = (score, b)
    order = ["Direct travelers", "Travel trade"]
    out = [best[a][1] for a in order if a in best]
    out += [v[1] for a, v in best.items() if a not in order]
    return out


def _extract_sources(section: Section) -> list[dict]:
    """Citations -> [{label, url}]. Requires a URL; supports one-line
    ('Label — desc. URL') and multi-line ('1. Label' / 'URL' / 'desc') styles."""
    out: list[dict] = []
    lines: list[str] = []
    for b in section.blocks:
        if b.items:
            lines.extend(b.items)
        elif b.text:
            lines.append(b.text)
    pending = ""
    for line in lines:
        m = _URL_RE.search(line)
        if not m:
            pending = line.strip()  # a label line for a following URL
            continue
        url = m.group(0).rstrip(".,);")
        label = (line[: m.start()]).strip().rstrip("—–-.) ").strip() or pending
        label = re.sub(r"^\d+[.)]\s*", "", label).strip()  # drop list numbering
        out.append({"label": label, "url": url})
        pending = ""
    return out


_REL_LINK_RE = re.compile(r"^[•·\-*\s]*(.+?)\s*(?:→|->)\s*(\S+)\s*$")


def _links_from_line(line: str) -> list[dict]:
    """Internal links from one footer line, in either house style:
    Markdown ``[Label](/url/)`` (Word hyperlink) or ``• Label → /url/``."""
    line = (line or "").strip()
    out: list[dict] = []
    md = _MD_LINK_RE.findall(line)  # Word hyperlinks -> [(label, url), ...]
    if md:
        for label, url in md:
            label = label.strip()
            if label and url.startswith(("/", "http")):
                out.append({"label": label, "url": url})
        return out
    m = _REL_LINK_RE.match(line)  # "Label → /url/"
    if m:
        label, url = m.group(1).strip(), m.group(2).strip()
        um = _MD_LINK_RE.search(url)  # url side may itself be a markdown link
        if um:
            url = um.group(2)
        label = _MD_LINK_RE.sub(r"\1", label).strip()
        if label and url.startswith(("/", "http")):
            out.append({"label": label, "url": url})
    return out


def _extract_related_links(section: Section) -> list[dict]:
    """Internal links from an 'Explore More' / 'Explore …' footer section."""
    out: list[dict] = []
    for b in section.blocks:
        lines = list(b.items) if b.items else ([b.text] if b.text else [])
        for line in lines:
            out.extend(_links_from_line(line))
    return out


def _extract_related_link_groups(section: Section) -> list[dict]:
    """Grouped 'Explore More' links: each in-section sub-heading (HEADING block)
    starts a new group whose links are the '• Label → /url/' lines under it.
    Returns ``[{"title": str, "links": [{"label","url"}, ...]}, ...]``."""
    groups: list[dict] = []
    current: dict | None = None
    for b in section.blocks:
        if b.type == BlockType.HEADING and b.text.strip():
            current = {"title": b.text.strip(), "links": []}
            groups.append(current)
            continue
        lines = list(b.items) if b.items else ([b.text] if b.text else [])
        for line in lines:
            for link in _links_from_line(line):
                if current is None:  # links before any sub-heading -> untitled group
                    current = {"title": "", "links": []}
                    groups.append(current)
                current["links"].append(link)
    return [g for g in groups if g["links"]]


_SMART_QUOTES = {
    "“": '"', "”": '"', "″": '"',
    "‘": "'", "’": "'", "′": "'",
}


def _normalize_quotes(text: str) -> str:
    """Word turns straight quotes into curly ones, which breaks pasted JSON.
    Fold the curly variants back so a doc-authored schema still parses."""
    for bad, good in _SMART_QUOTES.items():
        text = text.replace(bad, good)
    return text


def _balanced_json(text: str, start: int) -> str:
    """Return the {…}/[…] beginning at index `start`, matched by bracket depth
    (ignoring brackets inside strings). Empty string if unbalanced."""
    open_ch = text[start]
    close_ch = "}" if open_ch == "{" else "]"
    depth = 0
    in_str = False
    esc = False
    for i in range(start, len(text)):
        c = text[i]
        if in_str:
            if esc:
                esc = False
            elif c == "\\":
                esc = True
            elif c == '"':
                in_str = False
            continue
        if c == '"':
            in_str = True
        elif c == open_ch:
            depth += 1
        elif c == close_ch:
            depth -= 1
            if depth == 0:
                return text[start:i + 1]
    return ""


def _docx_full_text(path: str | Path) -> str:
    """Full text of a .docx, including text boxes / content controls that
    ``python-docx``'s ``.paragraphs`` skips (the WEBMASTER schema blocks live
    there). Reads word/document.xml directly and strips the tags."""
    import html
    import zipfile

    try:
        with zipfile.ZipFile(str(path)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
    except (KeyError, OSError, zipfile.BadZipFile):
        return ""
    xml = re.sub(r"</w:p>", "\n", xml)          # paragraph breaks
    xml = re.sub(r"<[^>]+>", "", xml)           # drop tags
    return html.unescape(xml)


def _extract_schema_jsonld(text: str) -> str:
    """Collect every schema.org block the uploader wrote INTO the document and
    normalize the two authoring styles into ONE ``@graph``:

    - a single ``{"@context":…,"@graph":[…]}`` block (Santa Cruz style), and
    - several separate labeled objects — ``TouristAttraction: {…}`` /
      ``FAQPage: {…}`` / ``BreadcrumbList: {…}`` (Baltra style).

    Curly quotes (Word) are folded so pasted JSON parses; label lines and the
    "WEBMASTER: paste this…" instructions are ignored (only valid JSON with an
    ``@context``/``@type`` is kept). Returns a compact JSON string, or "" when
    the document carries no schema. The engine never invents one.
    """
    text = _normalize_quotes(text or "")
    # Word wraps long values across lines; a literal newline inside a JSON string
    # is invalid, so fold line breaks/tabs to spaces before parsing.
    text = re.sub(r"[\r\n\t]+", " ", text)
    nodes: list = []
    covered: list[tuple[int, int]] = []
    for m in re.finditer(r"[{\[]", text):
        start = m.start()
        if any(a <= start < b for a, b in covered):
            continue  # inside an already-captured block
        block = _balanced_json(text, start)
        if not block or ("@context" not in block and "@type" not in block):
            continue
        try:
            obj = json.loads(block)
        except (ValueError, TypeError):
            continue
        covered.append((start, start + len(block)))
        items = obj if isinstance(obj, list) else [obj]
        for it in items:
            if isinstance(it, dict) and isinstance(it.get("@graph"), list):
                for n in it["@graph"]:
                    if isinstance(n, dict):
                        n.pop("@context", None)
                    nodes.append(n)
            elif isinstance(it, dict):
                it.pop("@context", None)
                nodes.append(it)
    if not nodes:
        return ""
    combined = {"@context": "https://schema.org", "@graph": nodes}
    return json.dumps(combined, ensure_ascii=False, separators=(",", ":"))


def _visual_heading_level(para) -> int:
    """Infer a heading level for docs that style headings by bold + font size
    rather than Word heading styles. 0 means 'not a heading'."""
    runs = [r for r in para.runs if (r.text or "").strip()]
    if not runs or not any(r.bold for r in runs):
        return 0
    text = para.text.strip()
    if len(text) > 140 or text.endswith((".", "!")):
        return 0
    sizes = [r.font.size.pt for r in runs if r.font.size is not None]
    size = max(sizes) if sizes else 0.0
    # A '?'-ending line is a heading only when it's clearly heading-sized and
    # short — a real title like "Baltra vs San Cristóbal — Which Airport?" —
    # never a body/FAQ question (which is set in the smaller body font).
    if text.endswith("?") and (size < 13.5 or len(text) > 80):
        return 0
    if size >= 18:
        return 1
    if size >= 13.5:
        return 2
    if size >= 11.5:
        return 3
    return 0


def _is_title_para(para) -> bool:
    style = (para.style.name if para.style else "").lower()
    return style in ("title", "heading 1") or _visual_heading_level(para) == 1


_SCI_NAME_RE = re.compile(r"\(([A-ZÁÉÍÓÚ][a-zé]+ [a-z]+)\)")


_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


def _split_off_button(paras: list[str]) -> tuple[list[str], dict | None]:
    """Pull a trailing button off a card's prose.

    Fires when the LAST paragraph is *only* link(s) (the internal-link notation
    sits on its own line) or an explicit '→ Label …/url' line. A paragraph that
    ends with a link but also carries prose is left alone — that link stays inline.
    """
    if not paras:
        return paras, None
    last = paras[-1].strip()
    body = re.sub(r"^(?:→|->)\s*", "", last)  # tolerate a leading arrow

    # The paragraph is purely link(s) (optionally "See also:" / "|" separated).
    links = _MD_LINK_RE.findall(body)
    residue = _MD_LINK_RE.sub("", body)
    residue = re.sub(r"(?i)see also:?|[\s,|·•]", "", residue).strip()
    if links and not residue:
        label, url = links[0]
        return paras[:-1], {"label": label.strip(), "url": url.strip()}

    # Explicit "→ Label /url" with a bare path/URL.
    if last.startswith(("→", "->")):
        arrow = re.search(r"(.+?)\s*(?:→|->)\s*(\S+)$", body)
        if arrow and arrow.group(2).startswith(("/", "http")):
            return paras[:-1], {"label": arrow.group(1).strip(), "url": arrow.group(2)}
    return paras, None


def _extract_cards(section: Section) -> tuple[str, list[tuple[str, list[str]]]]:
    """Split a prose section into (intro, [(h3_heading, paragraphs), …])."""
    intro: list[str] = []
    cards: list[tuple[str, list[str]]] = []
    name: str | None = None
    desc: list[str] = []
    for b in section.blocks:
        if b.type == BlockType.HEADING:
            if name is not None:
                cards.append((name, desc))
            name, desc = b.text, []
        elif b.text or b.items:
            chunk = "\n".join(b.items) if b.items else b.text
            (desc if name is not None else intro).append(chunk)
    if name is not None:
        cards.append((name, desc))
    return "\n\n".join(intro).strip(), cards


def _extract_visitor_sites_prose(section: Section) -> tuple[str, list[dict]]:
    """Visitor sites as H3 sub-headings + prose. Returns (section intro, sites)."""
    tl = section.title.lower()
    access = "Cruise-only" if "cruise" in tl else ("Land-based" if "land" in tl else "")
    intro, cards = _extract_cards(section)
    out: list[dict] = []
    for name, paras in cards:
        paras, button = _split_off_button(paras)
        row = {
            "site_name": name.strip(),
            "description": "\n\n".join(paras).strip(),
            "access_type": access,
        }
        if button:
            row["button_label"] = button["label"]
            row["button_url"] = button["url"]
        out.append(row)
    return intro, out


def _extract_where_to_see_prose(section: Section) -> tuple[str, list[dict]]:
    """'Where to See' H3 sites -> (intro prose, [{site, island, description}]).
    An 'Island — Site' heading splits into island + site."""
    intro, cards = _extract_cards(section)
    out: list[dict] = []
    for name, paras in cards:
        paras, _button = _split_off_button(paras)
        parts = re.split(r"\s+[—–]\s+", name.strip(), maxsplit=1)
        island, site = (parts[0].strip(), parts[1].strip()) if len(parts) == 2 else ("", name.strip())
        out.append({
            "site": site,
            "island": island,
            "description": "\n\n".join(paras).strip(),
        })
    return intro, out


def _section_prose_md(section: Section) -> str:
    """Render a section's blocks back to Markdown (H3 sub-headings + prose +
    callouts), in document order — used to keep a section's full narrative in a
    single WYSIWYG field."""
    parts: list[str] = []
    for b in section.blocks:
        if b.type == BlockType.HEADING:
            parts.append("### " + b.text.strip())
        elif b.items:
            parts.extend("- " + it for it in b.items)
        elif b.text:
            parts.append(b.text.strip())
    return "\n\n".join(p for p in parts if p).strip()


def _extract_wildlife(section: Section) -> tuple[str, list[dict]]:
    """Wildlife section -> (intro prose, species list).

    The H2 lead text (before the first H3) is the intro; each H3 sub-heading +
    its prose is a species, with an optional trailing '→' button.
    """
    intro: list[str] = []
    species: list[dict] = []
    name: str | None = None
    desc: list[str] = []

    def make(nm: str, paras: list[str]) -> dict:
        paras, button = _split_off_button(paras)
        text = "\n\n".join(paras).strip()
        sci = _SCI_NAME_RE.search(text)
        row = {
            "common_name": re.split(r"\s+[—–-]\s+", nm, maxsplit=1)[0].strip(),
            "scientific_name": sci.group(1) if sci else "",
            "description": text,
        }
        if button:
            row["button_label"] = button["label"]
            row["button_url"] = button["url"]
        return row

    for b in section.blocks:
        if b.type == BlockType.HEADING:
            if name and desc:
                species.append(make(name, desc))
            name, desc = b.text, []
        elif b.text or b.items:
            chunk = "\n".join(b.items) if b.items else b.text
            (desc if name else intro).append(chunk)
    if name and desc:
        species.append(make(name, desc))
    return "\n\n".join(intro).strip(), species
_KNOWN_META_KEYS = {
    "type",
    "page type",
    "template",
    "title",
    "slug",
    "destination",
    "duration",
    "price",
    "focus keyword",
    "keyword",
    "meta description",
    "description",
    "categories",
    "tags",
    "status",
    "ship",
    "region",
    "country",
    "best time",
    "featured image query",
    "featured image alt",
}


def read_docx(path: str | Path, *, page_type: str | None = None) -> Document:
    from docx import Document as DocxDocument  # imported lazily

    path = Path(path)
    docx = DocxDocument(str(path))

    # Page-specific schema.org block the author wrote INTO the document (in a
    # WEBMASTER text box that .paragraphs can't see, so read the raw XML). The
    # engine never generates schema — it publishes only what the doc provides.
    full_text = _docx_full_text(path)
    schema_block = _extract_schema_jsonld(full_text)

    # House "CMS Stage" docs use no heading styles and their own conventions;
    # route them to the dedicated adapter — UNLESS this is a wildlife species
    # page. Those docs carry the same "PUBLISHER HEADER BLOCK" marker but keep
    # their real content (facts, seasonality, subspecies, CTA) in data tables the
    # CMS adapter drops, so they must go through the normal table-aware path.
    from .cms import build_cms_document, looks_like_cms

    all_texts = [p.text for p in docx.paragraphs]
    is_species = _norm_page_type(page_type) in _SPECIES_PAGE_TYPES or _looks_like_species(full_text)
    is_informative = _norm_page_type(page_type) in _INFORMATIVE_PAGE_TYPES
    if not is_species and looks_like_cms([t for t in all_texts if t.strip()]):
        cms_doc = build_cms_document(all_texts, path.name)
        if schema_block:
            cms_doc.metadata["schema_jsonld"] = schema_block
        return cms_doc

    doc = Document(source_name=path.name, source_kind="docx")
    if schema_block:
        doc.metadata["schema_jsonld"] = schema_block

    # A real H1 / "Title"-styled line in the body wins. The core-properties title
    # is only a fallback — Word frequently leaves the generic "Word Document"
    # there, which must never become the page title.
    core_title = (docx.core_properties.title or "").strip()

    current = Section(title="", level=1, slug="_lead")
    # Document-order position of each block element, so an (unrecognized) table
    # can be attached to the section it actually sits under — not the last one.
    body_order = {el: i for i, el in enumerate(docx.element.body)}
    section_start_pos: list[tuple[int, Section]] = [(-1, current)]
    list_items: list[str] = []
    list_ordered = False
    list_pos = 1 << 30
    seen_body = False
    verify_warnings: list[str] = []
    geo_capturing = False       # inside a "[CARD START — GEO SUMMARY]…[CARD END]"
    geo_lines: list[str] = []
    # Reconstruct a markdown-ish body so the freeform builder can read any
    # directives an author typed directly into Word (e.g. "::: columns").
    raw_lines: list[str] = []

    def flush_list() -> None:
        nonlocal list_items, list_ordered
        if list_items:
            blk = ContentBlock(type=BlockType.LIST, ordered=list_ordered, items=list_items[:])
            blk.meta["_pos"] = list_pos
            current.blocks.append(blk)
            list_items = []
            list_ordered = False

    # Skip a cover-letter preamble: start at the document title (styled or a
    # bold, large-font line) when one exists.
    paras = list(docx.paragraphs)
    start = 0
    for i, p in enumerate(paras):
        txt = _HEADING_MARK.sub("", _p_text(p).strip()).strip()
        if txt and _is_title_para(p) and not _is_scaffold_title(txt):
            start = i
            break

    for para in paras[start:]:
        text = _clean_markers(_p_text(para)).strip()
        style = (para.style.name if para.style else "") or ""
        style_l = style.lower()
        cur_pos = body_order.get(para._p, 1 << 29)

        if not text:
            continue

        # Visual separator rule (────, ____) -> layout, not content.
        if _is_separator(text):
            continue

        # Strip literal house markers like "[H2] " from heading/paragraph text.
        text = _HEADING_MARK.sub("", text).strip()
        # Drop (and surface) editorial flags: "⚠️ WEBMASTER: Do not publish…".
        if _EDITORIAL_FLAG.match(text):
            verify_warnings.append("VERIFY: " + text.lstrip("⚠️ ").strip())
            continue
        # A stray "CTA button:" directive in the body is an instruction, not prose.
        if text.lower().startswith("cta button:"):
            continue
        # GEO citation card written as prose: capture the inner text as geo_answer,
        # drop the "[CARD START — GEO SUMMARY]" / "[CARD END]" markers.
        if _CARD_GEO_START_RE.match(text):
            geo_capturing = True
            continue
        if geo_capturing:
            if _CARD_END_RE.match(text):
                geo_capturing = False
            else:
                geo_lines.append(text)
            continue
        if _CARD_END_RE.match(text):
            continue
        # CMS-Stage annotations / banners / publisher strip — never body content.
        if (_BRACKET_ANNOT_RE.match(text) or _STAGE_SCAFFOLD_RE.match(text)
                or _PUBLISHER_LINE_RE.search(text)):
            continue
        # GEO/AIO snippet scaffolding some docs write as prose (not a table):
        # drop the instruction lines and unwrap the "Text: <answer>" prefix so the
        # publishable answer stays as the lead paragraph.
        if _GEO_SCAFFOLD_RE.match(text):
            continue
        if len(text) > 40:
            text = re.sub(r"^Text:\s+", "", text)

        # Heading from a heading/title style, or visually (bold + larger font for
        # docs that don't use Word heading styles).
        if style_l.startswith("heading"):
            level = _heading_level(style_l)
        elif style_l == "title":
            level = 1
        else:
            level = _visual_heading_level(para)

        if level:
            flush_list()
            if level == 1 and not doc.title and not seen_body:
                # A CMS scaffold label ("… Page", the domain) is not the title.
                if _is_scaffold_title(text):
                    continue
                doc.title = text
                continue
            if level >= 3:
                # Subheadings nest inside the current section (e.g. FAQ items).
                h_blk = ContentBlock(type=BlockType.HEADING, text=text, level=level)
                h_blk.meta["_pos"] = cur_pos
                current.blocks.append(h_blk)
                raw_lines.append("#" * level + " " + text)
                seen_body = True
                continue
            if current.blocks or current.title:
                doc.sections.append(current)
            current = Section(title=text, level=level, slug=section_slug(text))
            section_start_pos.append((body_order.get(para._p, -1), current))
            raw_lines.extend(["", "#" * max(2, level) + " " + text])
            seen_body = True
            continue

        # A byline anywhere in the lead (before the first H2) -> author metadata,
        # even when it follows the intro paragraph. `not current.title` means we
        # are still in the lead section, so a mid-article "By …" stays as prose.
        if not doc.metadata.get("author") and not current.title and _looks_like_byline(text):
            doc.metadata["author"] = _clean_byline(text)
            continue

        # Leading "Key: value" lines before any prose -> metadata.
        if not seen_body and not current.blocks:
            m = _META_RE.match(text)
            if m and m.group(1).strip().lower() in _KNOWN_META_KEYS:
                key = m.group(1).strip().lower().replace(" ", "_")
                doc.metadata[key] = m.group(2).strip()
                continue

        seen_body = True

        # List styles.
        if "list bullet" in style_l or "list number" in style_l or style_l.startswith("list"):
            if not list_items:
                list_pos = cur_pos
            list_ordered = "number" in style_l
            list_items.append(text)
            raw_lines.append(("1. " if list_ordered else "- ") + text)
            continue
        flush_list()

        if style_l == "quote" or style_l == "intense quote":
            q_blk = ContentBlock(type=BlockType.QUOTE, text=text)
            q_blk.meta["_pos"] = cur_pos
            current.blocks.append(q_blk)
            raw_lines.extend(["", "> " + text])
            continue

        p_blk = ContentBlock(type=BlockType.PARAGRAPH, text=text)
        p_blk.meta["_pos"] = cur_pos
        current.blocks.append(p_blk)
        raw_lines.extend(["", text])

    flush_list()
    if current.blocks or current.title:
        doc.sections.append(current)
    doc.raw_body = "\n".join(raw_lines).strip()

    # GEO citation-card prose captured above -> geo_answer (unless a table already
    # supplied one). A trailing "Source: …" attribution is dropped from the
    # concise answer.
    if geo_lines and not doc.metadata.get("geo_answer"):
        geo = " ".join(g.strip() for g in geo_lines if g.strip()).strip()
        geo = re.split(r"\s*Source:\s", geo, maxsplit=1)[0].strip()
        if geo:
            doc.metadata["geo_answer"] = geo

    # Tables become table blocks appended to the lead/last section. Internal
    # instruction / explainer tables (green boxes) are skipped — not content.
    season_cands: list[list[dict]] = []
    for table in docx.tables:
        rows = [[_clean_markers(_cell_text(cell)) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        geo = _extract_geo_answer(rows)
        if geo:
            # A box that is a webmaster INSTRUCTION about the Quick Answer ("This
            # block is your GEO snippet… Place it as the VERY FIRST element") is
            # not the answer — drop it instead of publishing it as geo_answer.
            if _GEO_INSTRUCTION_RE.search(geo[:200]):
                continue
            doc.metadata.setdefault("geo_answer", geo)
            continue  # the GEO block is scaffolding, not body content
        if _is_geo_instruction(rows):
            continue  # webmaster "how to place the Quick Answer" box -> drop
        if _is_instruction_table(rows):
            continue
        # Webmaster / SEO scaffold KV tables (PUBLISH URL, Title tag, Anchor
        # Text…) that species docs pack into 2-col tables — not content, but a
        # header block often carries the Author / Byline, so mine that first.
        if _is_scaffold_kv_table(rows):
            if not doc.metadata.get("author"):
                aut = _author_from_kv(rows)
                if aut:
                    doc.metadata["author"] = aut
            continue
        # Single-cell boxes: scaffolding/placeholders are dropped; a prose box is
        # a pull-quote/callout, not a full-width data table.
        if _ncols(rows) == 1 and len(rows) == 1:
            one = rows[0][0].strip()
            if not one:
                continue
            if one.startswith("[") or _SCAFFOLD_RE.match(one) or _looks_like_jsonld(one):
                continue  # placeholder, PUBLISH/SLUG scaffold, or a JSON-LD block
            # A "CTA: <audience> | <company>" pull-out is a real CTA, not prose.
            if re.match(r"^\s*CTA\b", one, re.IGNORECASE) or _is_cta_table(rows):
                blocks = _extract_cta_blocks(rows)
                if blocks:
                    doc.metadata.setdefault("cta_blocks", []).extend(blocks)
                    continue
            blk = ContentBlock(type=BlockType.PARAGRAPH, text=one)
            blk.meta["_pos"] = body_order.get(table._tbl, 1 << 30)
            _table_target(table, body_order, section_start_pos, doc, current).blocks.append(blk)
            continue
        if not is_informative and _is_visitor_sites_table(rows):
            sites = _extract_visitor_sites(rows)
            if sites:
                doc.metadata.setdefault("visitor_sites", []).extend(sites)
                continue
        if _is_sources_table(rows):
            srcs = _extract_sources_table(rows)
            if srcs:
                doc.metadata.setdefault("sources", []).extend(srcs)
                continue
        if not is_informative and _is_wildlife_calendar_table(rows):
            cal = _extract_wildlife_calendar(rows)
            if cal:
                doc.metadata.setdefault("wildlife_calendar", []).extend(cal)
                continue
        # Per-island variants (giant tortoise shell types, marine iguana races).
        if not is_informative and _is_subspecies_table(rows):
            sub = _extract_subspecies(rows)
            if sub:
                doc.metadata.setdefault("subspecies", []).extend(sub)
                continue
        # Viewing calendar / breeding cycle (period + optional status + notes).
        # Several may exist (e.g. a "Period | Event" plus a "Month | Status |
        # Notes"); keep the richest and remember the rest as candidates.
        if not is_informative and _is_seasonality_table(rows):
            season = _extract_seasonality(rows)
            if season:
                season_cands.append(season)
                continue
        if not is_informative and _is_quick_facts_table(rows):
            facts = _extract_quick_facts(rows)
            if facts:
                doc.metadata.setdefault("quick_facts", []).extend(facts)
                continue
        if _is_cta_table(rows):
            blocks = _extract_cta_blocks(rows)
            if blocks:
                doc.metadata.setdefault("cta_blocks", []).extend(blocks)
                continue
        # Attach to the section active at this table's document position, so a
        # fees/comparison table lands under its own heading (not dumped into the
        # last section, e.g. "Sources").
        target = _table_target(table, body_order, section_start_pos, doc, current)
        tbl_blk = ContentBlock(type=BlockType.TABLE, rows=rows)
        tbl_blk.meta["_pos"] = body_order.get(table._tbl, 1 << 30)
        target.blocks.append(tbl_blk)

    # Restore document order within each section. Stray single-cell tables
    # (pull-quote/stat callouts) are appended after the paragraph pass, so sort
    # every section's blocks by their source position — otherwise the callouts
    # pile up at the end of the section instead of sitting where they belong.
    for s in doc.sections:
        s.blocks.sort(key=lambda b: b.meta.get("_pos", 1 << 30))

    # Plain-paragraph CTAs (waved-albatross style): some docs write the two
    # formal CTAs as body paragraphs ("BOOK WITH VOYAGERS… —", "TRAVEL INDUSTRY
    # PARTNERS… —") rather than a CTA table. Lift them into cta_blocks and drop
    # them from the section so they render as CTA cards instead of leaking into
    # the Plan-Your-Visit intro prose.
    for s in doc.sections:
        kept = []
        for blk in s.blocks:
            if blk.type == BlockType.PARAGRAPH and _looks_like_plain_cta(blk.text or ""):
                cta = _cta_from_paragraph(blk.text)
                if cta:
                    doc.metadata.setdefault("cta_blocks", []).append(cta)
                    continue
            kept.append(blk)
        s.blocks = kept

    # Collapse the raw CTA candidates to the two canonical CTAs (drops any
    # scaffolding that leaked through a mixed green box).
    if doc.metadata.get("cta_blocks"):
        doc.metadata["cta_blocks"] = _dedupe_ctas(doc.metadata["cta_blocks"])

    # Seasonality: pick the richest calendar (most rows, tie-break: has a
    # status/label column) when a doc carries more than one.
    if season_cands:
        best_season = max(
            season_cands,
            key=lambda c: (len(c), sum(1 for r in c if r.get("label"))),
        )
        doc.metadata["seasonality"] = best_season

    # Dedicated species facts (scientific/common name, IUCN status, population)
    # mined from the "At a Glance" facts table; endemism from the body text.
    if doc.metadata.get("quick_facts"):
        for field, value in _extract_species_facts(doc.metadata["quick_facts"]).items():
            doc.metadata.setdefault(field, value)
    if re.search(r"\bendemic\b", full_text, re.IGNORECASE):
        doc.metadata.setdefault("endemic", True)

    # Citations / sources section -> structured metadata ('sources',
    # 'sources-citations', 'sources & citations'…).
    for s in doc.sections:
        if s.slug.startswith("sources") or s.title.lower().startswith("sources"):
            srcs = _extract_sources(s)
            if srcs:
                doc.metadata.setdefault("sources", []).extend(srcs)
            break

    # "Explore More" / "Explore <X>" footer -> related internal links (flat list +
    # grouped by sub-heading). Once mined, drop the section so it doesn't ALSO
    # render as a Feature Section — it's the SEO footer, not body content.
    for s in doc.sections:
        if "explore" in s.slug or "footer" in s.slug or s.title.lower().startswith("explore"):
            related = _extract_related_links(s)
            if related:
                doc.metadata["related_links"] = related
            groups = _extract_related_link_groups(s)
            if groups:
                doc.metadata["related_link_groups"] = groups
            if related or groups:
                doc.sections.remove(s)
            break

    # "At a Glance" / "Quick Facts" section -> the quick-facts heading + intro.
    for s in doc.sections:
        tl = s.title.lower()
        if "at a glance" in tl or "quick facts" in tl or "glance" in s.slug:
            doc.metadata["quick_facts_title"] = s.title
            intro = "\n\n".join(
                ("\n".join(b.items) if b.items else b.text)
                for b in s.blocks
                if b.type in (BlockType.PARAGRAPH, BlockType.LIST) and (b.text or b.items)
            ).strip()
            if intro:
                doc.metadata["quick_facts_intro"] = intro
            break

    # A "Wildlife" section with H3 species sub-headings -> title + intro + species.
    # Match "wildlife" as a whole word anywhere in the heading so free-form
    # titles ("The Wildlife", "Christmas Iguanas and Other Wildlife") still
    # qualify. When several sections mention wildlife, keep the one that yields
    # the most species rows (so we don't stop on a passing mention that has none).
    best: tuple[int, Section, str, list[dict]] | None = None
    for s in doc.sections:
        if not re.search(r"\bwildlife\b", s.slug) and not re.search(
            r"\bwildlife\b", s.title.lower()
        ):
            continue
        intro, wildlife = _extract_wildlife(s)
        if not wildlife:
            continue
        if best is None or len(wildlife) > best[0]:
            best = (len(wildlife), s, intro, wildlife)
    if best is not None:
        _, s, intro, wildlife = best
        doc.metadata["wildlife_title"] = s.title
        if intro:
            doc.metadata["wildlife_intro"] = intro
        doc.metadata["wildlife"] = wildlife

    # "Visitor Sites" sections -> the title + prose-based sites (H3 sub-headings).
    # (A table-based extraction may already have filled visitor_sites upstream.)
    for s in doc.sections:
        if s.slug.startswith("visitor-sites") or s.title.lower().startswith("visitor sites"):
            doc.metadata.setdefault(
                "visitor_sites_title", re.split(r"\s+[—–-]\s+", s.title, maxsplit=1)[0].strip()
            )
            intro, prose_sites = _extract_visitor_sites_prose(s)
            # Land-Based and Cruise-Only sections keep their own lead text.
            if intro:
                key = (
                    "visitor_sites_intro_cruise"
                    if "cruise" in s.title.lower()
                    else "visitor_sites_intro"
                )
                doc.metadata.setdefault(key, intro)
            if prose_sites:
                doc.metadata.setdefault("visitor_sites", []).extend(prose_sites)

    # Wildlife species pages: pull the prose "Where to See" and "Subspecies"
    # sections into their dedicated repeaters/fields (gated to species docs so
    # island pages are untouched). Both are then removed from doc.sections so
    # they don't ALSO render as Feature Sections.
    if is_species:
        for s in list(doc.sections):
            tl = s.title.lower()
            if s.slug == "where_to_see" or tl.startswith(("where to see", "where and how", "where to find")):
                intro, sites = _extract_where_to_see_prose(s)
                if sites:
                    doc.metadata.setdefault("where_to_see_title", s.title)
                    if intro:
                        doc.metadata.setdefault("where_to_see_intro", intro)
                    doc.metadata["where_to_see"] = sites
                    doc.sections.remove(s)
                break
        for s in list(doc.sections):
            if "subspecies" in s.slug or "subspecies" in s.title.lower():
                doc.metadata.setdefault("subspecies_title", s.title)
                intro_md = _section_prose_md(s)  # intro + the H3 island narratives
                if intro_md:
                    doc.metadata.setdefault("subspecies_intro", intro_md)
                doc.sections.remove(s)
                break

    if verify_warnings:
        doc.metadata["_ingest_warnings"] = [
            *verify_warnings,
            "This doc has unresolved VERIFY flags — review before publishing live.",
        ]

    if not doc.title:
        if core_title and core_title.lower() != "word document":
            doc.title = core_title
        else:
            doc.title = path.stem.replace("_", " ").replace("-", " ").title()
    return doc


def _heading_level(style_name: str) -> int:
    m = re.search(r"(\d)", style_name)
    return int(m.group(1)) if m else 2
